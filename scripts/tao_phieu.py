# -*- coding: utf-8 -*-
"""Đọc file .json trong HANG_CHO (cục bộ + Google Drive) -> tạo Phiếu YCPT Hóa học và/hoặc Sinh học (.docx).

Chạy:  python tao_phieu.py [duong_dan_file_excel]
"""
import datetime
import os
import re
import shutil
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pyc_common import classify, norm, paginate, split_chi_tieu, UNKNOWN  # noqa: E402
import pyc_fill  # noqa: E402

ROOT = pyc_fill.ROOT
OUT_DIR = ROOT  # kết quả: <thư mục>\\HOA_HOC và <thư mục>\\VI_SINH

INFO_KEYS = {
    "khách hàng": "khach_hang",
    "địa chỉ": "dia_chi",
    "điện thoại": "dien_thoai",
    "email": "email",
    "số hh (phiếu hóa)": "so_hh",
    "số vs (phiếu sinh)": "so_vs",
    "tài liệu kèm theo": "tai_lieu_kem",
    "các yêu cầu khác": "yeu_cau_khac",
    "số mẫu/trang hóa học": "per_hoa",
    "số mẫu/trang sinh học": "per_sinh",
    "tên file kết quả": "ten_file",
}


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return norm(v)


def read_input(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["THONG_TIN"]
    info = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        if not row or row[0] is None:
            continue
        k = norm(row[0]).lower().rstrip(":")
        if k in INFO_KEYS:
            info[INFO_KEYS[k]] = _cell(row[1] if len(row) > 1 else "")
    ws = wb["DANH_SACH_MAU"]
    header = [norm(c.value).lower() for c in ws[1]]

    def col(name):
        for i, h in enumerate(header):
            if h.startswith(name):
                return i
        raise KeyError(f"Thiếu cột '{name}' trong sheet DANH_SACH_MAU")

    ci = {k: col(k) for k in ["stt", "số nhận diện", "mô tả", "khối lượng", "nền mẫu", "chỉ tiêu"]}
    samples = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        mota = _cell(r[ci["mô tả"]])
        ct = _cell(r[ci["chỉ tiêu"]]) if r[ci["chỉ tiêu"]] is not None else ""
        if not mota and not ct:
            continue
        samples.append({
            "stt_goc": _cell(r[ci["stt"]]),
            "ma": _cell(r[ci["số nhận diện"]]),
            "mota": mota,
            "kl": _cell(r[ci["khối lượng"]]),
            "nen": _cell(r[ci["nền mẫu"]]).lower(),
            "chi_tieu_raw": str(r[ci["chỉ tiêu"]] or ""),
        })
    return info, samples


def validate(info, samples):
    errs = []
    if not info.get("khach_hang"):
        errs.append("Chưa nhập Khách hàng (sheet THONG_TIN)")
    if not samples:
        errs.append("Chưa có mẫu nào trong sheet DANH_SACH_MAU")
    for i, s in enumerate(samples, 1):
        if not s["mota"]:
            errs.append(f"Dòng mẫu {i}: thiếu Mô tả mẫu")
        if not s["kl"]:
            errs.append(f"Dòng mẫu {i} ({s['mota']}): thiếu Khối lượng")
        if not split_chi_tieu(s["chi_tieu_raw"]):
            errs.append(f"Dòng mẫu {i} ({s['mota']}): thiếu Chỉ tiêu")
    return errs


def relaunder_linux(path):
    """Chỉ dùng khi test trên Linux; trên Windows file .bat sẽ cho Word mở & lưu lại."""
    soffice = shutil.which("soffice")
    if not soffice:
        return False
    tmp = tempfile.mkdtemp()
    os.system(f'"{soffice}" --headless --convert-to docx:"MS Word 2007 XML" --outdir "{tmp}" "{path}" >/dev/null 2>&1')
    out = os.path.join(tmp, os.path.basename(path))
    if os.path.exists(out):
        shutil.move(out, path)
        return True
    return False


def is_water(s):
    return s["nen"].startswith("nước") or s["nen"].startswith("nuoc")


def make_form(kind, info, samples, per_page, out_path, warnings):
    """kind = 'H' | 'S'. samples đã lọc chỉ còn chỉ tiêu đúng loại."""
    for i, s in enumerate(samples, 1):
        s["stt"] = str(i)
    if kind == "H":
        pages = paginate(samples, per_page, max_ct=pyc_fill.MAX_HOA_CT)
    else:
        # Mẫu nước và mẫu thường không chung 1 trang (đơn vị khác nhau) — vẫn giữ thứ tự
        pages, run, flag = [], [], None
        for s in samples:
            f = is_water(s)
            if run and f != flag:
                pages += paginate(run, per_page)
                run = []
            run.append(s)
            flag = f
        pages += paginate(run, per_page)
    tmpdir = tempfile.mkdtemp()
    files = []
    for pi, page in enumerate(pages, 1):
        pinfo = dict(info)
        pinfo["so_phieu"] = info.get("so_hh" if kind == "H" else "so_vs", "")
        pinfo["so_luong_mau"] = str(len(samples))
        pinfo["trang_so"], pinfo["tong_trang"] = pi, len(pages)
        f = os.path.join(tmpdir, f"p{pi}.docx")
        if kind == "H":
            pyc_fill.build_hoa_page(f, pinfo, page, warnings)
        else:
            nuoc = is_water(page[0])
            pyc_fill.build_sinh_page(f, pinfo, page, warnings, nen_nuoc=nuoc)
        files.append(f)
    if len(files) == 1:
        shutil.copy(files[0], out_path)
    else:
        pyc_fill.compose_pages(files, out_path)
    return [len(p) for p in pages]


def safe_name(s):
    s = re.sub(r'[\\/:*?"<>|]+', "_", s).strip()
    if len(s) > 45:
        s = s[:45].rsplit(" ", 1)[0]
    return s or "phieu"


LOG_FILE = os.path.join(ROOT, "SO_THEO_DOI_PYC.xlsx")
LOG_PHIEU_HDR = ["Ngày tạo", "Loại", "Số phiếu", "Khách hàng", "Số mẫu", "Mã mẫu (từ – đến)",
                 "Chỉ tiêu", "Tài liệu kèm theo", "Nguồn dữ liệu", "File phiếu"]
LOG_MAU_HDR = ["Ngày tạo", "Loại", "Số phiếu", "Số nhận diện mẫu", "Mô tả mẫu", "KL (g)",
               "Chỉ tiêu", "Khách hàng"]


def read_json(path):
    """Phiếu gửi từ điện thoại (hàng chờ): {"thong_tin": {...}, "mau": [{...}]}"""
    import json
    with open(path, encoding="utf-8") as f:
        d = json.load(f)
    info = {k: _cell(v) for k, v in d.get("thong_tin", {}).items()}
    samples = []
    for m in d.get("mau", []):
        samples.append({"stt_goc": _cell(m.get("stt", "")), "ma": _cell(m.get("ma", "")),
                        "ma_hh": _cell(m.get("ma_hh", "")), "ma_vs": _cell(m.get("ma_vs", "")),
                        "mota": _cell(m.get("mota", "")), "kl": _cell(m.get("kl", "")),
                        "nen": _cell(m.get("nen", "")).lower(), "chi_tieu_raw": str(m.get("chi_tieu", ""))})
    return info, samples


def unique_path(path):
    if not os.path.exists(path):
        return path
    b, e = os.path.splitext(path)
    i = 2
    while os.path.exists(f"{b}_v{i}{e}"):
        i += 1
    return f"{b}_v{i}{e}"


def write_log(rows_phieu, rows_mau, warnings_out):
    """Ghi sổ theo dõi. Nếu file đang mở trong Excel -> ghi tạm sang file chờ, lần sau gộp."""
    pending = os.path.join(ROOT, "scripts", "_so_theo_doi_cho_ghi.json")
    import json
    queue = {"phieu": [], "mau": []}
    if os.path.exists(pending):
        try:
            with open(pending, encoding="utf-8") as f:
                queue = json.load(f)
        except Exception:
            pass
    queue["phieu"] += rows_phieu
    queue["mau"] += rows_mau
    try:
        if os.path.exists(LOG_FILE):
            wb = openpyxl.load_workbook(LOG_FILE)
        else:
            from openpyxl.styles import Font, PatternFill
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "PHIEU"
            ws.append(LOG_PHIEU_HDR)
            ws2 = wb.create_sheet("MAU")
            ws2.append(LOG_MAU_HDR)
            for w in (ws, ws2):
                for c in w[1]:
                    c.font = Font(bold=True, color="FFFFFF")
                    c.fill = PatternFill("solid", fgColor="1F4E79")
                w.freeze_panes = "A2"
                w.auto_filter.ref = w.dimensions
            for col, wd in zip("ABCDEFGHIJ", [16, 9, 13, 40, 8, 20, 45, 40, 30, 60]):
                ws.column_dimensions[col].width = wd
            for col, wd in zip("ABCDEFGH", [16, 9, 13, 14, 45, 8, 40, 40]):
                ws2.column_dimensions[col].width = wd
        for r in queue["phieu"]:
            wb["PHIEU"].append(r)
        for r in queue["mau"]:
            wb["MAU"].append(r)
        for name in ("PHIEU", "MAU"):
            wb[name].auto_filter.ref = wb[name].dimensions
        wb.save(LOG_FILE)
        if os.path.exists(pending):
            os.remove(pending)
    except PermissionError:
        with open(pending, "w", encoding="utf-8") as f:
            json.dump(queue, f, ensure_ascii=False)
        warnings_out.append("Sổ theo dõi đang mở trong Excel — đã ghi tạm, sẽ tự gộp vào sổ ở lần chạy sau")


def code_range(lst):
    codes = [x["ma"] for x in lst if x.get("ma")]
    if not codes:
        return ""
    return codes[0] if len(codes) == 1 else f"{codes[0]} – {codes[-1]}"


def existing_numbers():
    try:
        wb = openpyxl.load_workbook(LOG_FILE, read_only=True)
        return {(r[1], str(r[2])): r[0] for r in wb["PHIEU"].iter_rows(min_row=2, values_only=True) if r and r[2]}
    except Exception:
        return {}


CODE_RE = re.compile(r"^L\d{3,5}$")


def check_codes(hoa, sinh):
    """Mã Laabb: đúng dạng, không trùng trong phiếu, không trùng giữa Hóa/Sinh, không trùng sổ (cùng tháng)."""
    out, seen = [], {}
    for loai, lst in (("Hóa", hoa), ("Sinh", sinh)):
        for x in lst:
            c = x.get("ma", "")
            if not c:
                continue
            if not CODE_RE.match(c):
                out.append(f"Mã '{c}' ({x['mota']}) sai dạng Laabb")
            if c in seen:
                out.append(f"Mã {c} dùng 2 lần ({seen[c]} và {loai}: {x['mota']})")
            seen[c] = f"{loai}: {x['mota']}"
    try:
        thang = datetime.datetime.now().strftime("/%m/%Y")
        wb = openpyxl.load_workbook(LOG_FILE, read_only=True)
        for r in wb["MAU"].iter_rows(min_row=2, values_only=True):
            if r and r[3] in seen and str(r[0]).endswith(thang):
                out.append(f"Mã {r[3]} đã có trong sổ theo dõi (phiếu {r[2]}, {r[0]}) — kiểm tra trùng")
    except Exception:
        pass
    return out


def process(info, samples, src_desc, src_path=None):
    """Tạo phiếu cho 1 bộ dữ liệu. Trả về (mã lỗi, danh sách file đã tạo)."""
    errs = validate(info, samples)
    if errs:
        print(f"\n*** [{src_desc}] CHƯA TẠO PHIẾU — cần bổ sung:")
        for e in errs:
            print("   -", e)
        return 2, []

    hoa, sinh, code_warn = [], [], []
    UNKNOWN.clear()
    for s in samples:
        h_ct, s_ct = [], []
        for c in split_chi_tieu(s["chi_tieu_raw"]):
            kind, name = classify(c)
            lst_ = s_ct if kind == "S" else h_ct
            if name not in lst_:
                lst_.append(name)
        # Mã Laabb: phiếu Hóa và phiếu Sinh dùng 2 mã KHÁC nhau (02.11-ST)
        ma, ma_hh, ma_vs = s.get("ma", ""), s.get("ma_hh", ""), s.get("ma_vs", "")
        if h_ct and s_ct:
            if ma and not (ma_hh or ma_vs):
                code_warn.append(f"Mẫu '{s['mota']}' có cả Hóa và Sinh nhưng chỉ có 1 mã ({ma}) — "
                                 f"đã đặt cho phiếu Hóa, mã Sinh để trống (2 phiếu phải là 2 mã khác nhau)")
            ma_hh = ma_hh or ma
            if ma_hh and ma_vs and ma_hh == ma_vs:
                code_warn.append(f"Mẫu '{s['mota']}': mã Hóa và mã Sinh trùng nhau ({ma_hh}) — phải là 2 số khác nhau")
        else:
            ma_hh = ma_hh or ma
            ma_vs = ma_vs or ma
        if h_ct:
            hoa.append({**s, "ma": ma_hh, "chi_tieu": h_ct})
        if s_ct:
            sinh.append({**s, "ma": ma_vs, "chi_tieu": s_ct})
    code_warn += check_codes(hoa, sinh)

    now = datetime.datetime.now()
    ym = (now.strftime("%Y"), now.strftime("%Y-%m"))
    base = safe_name(info.get("ten_file") or info["khach_hang"])
    per_hoa = int(info.get("per_hoa") or 5)
    per_sinh = int(info.get("per_sinh") or 6)
    created, report, log_p, log_m, extra_warn = [], [], [], [], []
    seen = existing_numbers()
    for kind, lst, per, tag, so in (("H", hoa, per_hoa, "HOA_HOC", info.get("so_hh")),
                                    ("S", sinh, per_sinh, "VI_SINH", info.get("so_vs"))):
        if not lst:
            continue
        warnings = []
        loai = "Hóa học" if kind == "H" else "Vi sinh"
        if not so:
            warnings.append(f"Số {'HH' if kind == 'H' else 'VS'} để trống — điền tay sau khi xác nhận số thứ tự trong ngày")
        elif (loai, so) in seen:
            warnings.append(f"Số {so} ĐÃ có phiếu tạo lúc {seen[(loai, so)]} — bản mới lưu thêm hậu tố _v2; kiểm tra có trùng số không")
        sub_dir = os.path.join(OUT_DIR, tag, ym[0], ym[1])
        os.makedirs(sub_dir, exist_ok=True)
        lead = safe_name(so) if so else "CHUA_SO_" + now.strftime("%d%m_%H%M")
        out = unique_path(os.path.join(sub_dir, f"{lead}_{base}.docx"))
        sizes = make_form(kind, info, [dict(x) for x in lst], per, out, warnings)
        if sys.platform != "win32" and len(sizes) > 1 and not os.environ.get("NO_RELAUNDER"):
            relaunder_linux(out)
        created.append(f"{out}|{len(sizes) + 1}")
        cts = []
        for x in lst:
            for c in x["chi_tieu"]:
                if c not in cts:
                    cts.append(c)
        report.append((tag, len(lst), sizes, cts, warnings, out))
        loai = "Hóa học" if kind == "H" else "Vi sinh"
        log_p.append([now.strftime("%d/%m/%Y %H:%M"), loai, so or "(chưa có)", info["khach_hang"], len(lst),
                      code_range(lst), "; ".join(cts), info.get("tai_lieu_kem", ""), src_desc,
                      os.path.relpath(out, ROOT)])
        for x in lst:
            log_m.append([now.strftime("%d/%m/%Y"), loai, so or "(chưa có)", x.get("ma", ""), x["mota"],
                          x.get("kl", ""), "; ".join(x["chi_tieu"]), info["khach_hang"]])

    # Lưu bản dữ liệu đầu vào cạnh phiếu để sau này sửa/tạo lại
    if src_path and os.path.exists(src_path):
        arch = os.path.join(OUT_DIR, "DU_LIEU_NHAP", ym[0], ym[1])
        os.makedirs(arch, exist_ok=True)
        so_all = "_".join(x for x in (info.get("so_hh"), info.get("so_vs")) if x) or "CHUA_SO_" + now.strftime("%d%m_%H%M")
        ext = os.path.splitext(src_path)[1]
        shutil.copy2(src_path, unique_path(os.path.join(arch, f"{safe_name(so_all)}_{base}{ext}")))

    write_log(log_p, log_m, extra_warn)

    print(f"\n================ KẾT QUẢ: {src_desc} ================")
    for tag, n, sizes, cts, warns, out in report:
        print(f"\n[{tag}] {n} mẫu, {len(sizes)} trang (mẫu/trang: {sizes})")
        print("   Chỉ tiêu:", "; ".join(cts))
        for w in warns:
            print("   ! ", w)
        print("   ->", os.path.relpath(out, ROOT))
    for w in extra_warn + code_warn:
        print("   ! ", w)
    for loai, t in sorted(UNKNOWN):
        print(f"   !  Chỉ tiêu {loai} '{t}' chưa có trong danh mục QĐ 1586 — kiểm tra tên"
              f" (hoặc thêm vào scripts/danh_muc_chi_tieu.json)")
    return 0, created


def _cfg(key):
    cfg = os.path.join(ROOT, "scripts", "cau_hinh.txt")
    vals = []
    if os.path.exists(cfg):
        for line in open(cfg, encoding="utf-8-sig"):
            if line.strip().upper().startswith(key.upper() + "="):
                v = line.split("=", 1)[1].strip()
                if v:
                    vals.append(v)
    return vals


def drive_dir():
    """Thư mục PYC_PT trên Google Drive (máy tính). Cấu hình: DRIVE=<đường dẫn>; mặc định tự dò ?:\\My Drive\\PYC_PT."""
    for v in _cfg("DRIVE"):
        if os.path.isdir(v):
            return v
    if sys.platform == "win32":
        for L in "GHIJKLMNOPQRSTUVWXYZEF":
            for name in ("My Drive", "Drive của tôi"):
                p = f"{L}:\\{name}\\PYC_PT"
                if os.path.isdir(p):
                    return p
    return None


def queue_dirs():
    """Hàng chờ: HANG_CHO cục bộ + HANG_CHO trên Google Drive (điện thoại/Claude bỏ file vào đây)."""
    dirs = _cfg("HANG_CHO") or [os.path.join(ROOT, "HANG_CHO")]
    d = drive_dir()
    if d:
        dirs.append(os.path.join(d, "HANG_CHO"))
    return [os.path.abspath(x) for x in dirs if os.path.isdir(x)]


TOOL_FILES = ["scripts/tao_phieu.py", "scripts/pyc_common.py", "scripts/pyc_fill.py",
              "scripts/danh_muc_chi_tieu.json",
              "mau_goc/PYC_PT_HOA_HOC_2025_11_17.docx", "mau_goc/PYC_PT_SINH_HOC_2025_11_17.docx"]


def update_tool_zip():
    """Đóng gói lại PYC_TOOL.zip khi script/mẫu thay đổi; chép thêm lên Drive để dùng từ điện thoại."""
    import zipfile
    z = os.path.join(ROOT, "PYC_TOOL.zip")
    src = [os.path.join(ROOT, f) for f in TOOL_FILES if os.path.exists(os.path.join(ROOT, f))]
    newest = max(os.path.getmtime(f) for f in src)
    if not os.path.exists(z) or os.path.getmtime(z) < newest:
        with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
            for f in src:
                zf.write(f, os.path.relpath(f, ROOT).replace("\\", "/"))
        print("   (đã cập nhật PYC_TOOL.zip)")
    d = drive_dir()
    if d:
        dz = os.path.join(d, "PYC_TOOL.zip")
        if not os.path.exists(dz) or os.path.getmtime(dz) < os.path.getmtime(z):
            shutil.copy2(z, dz)


def main(args):
    jobs = []
    if args:
        jobs = [os.path.abspath(a) for a in args]
    else:
        qs = queue_dirs()
        for q in qs:
            for f in sorted(os.listdir(q)):
                if f.lower().endswith((".json", ".xlsx")) and not f.startswith("~$"):
                    jobs.append(os.path.join(q, f))
        if jobs:
            print(f"Có {len(jobs)} phiếu trong hàng chờ ({'; '.join(qs)}) — xử lý trước.")
        default = os.path.join(ROOT, "NHAP_MAU.xlsx")  # tùy chọn, có thể xóa
        if os.path.exists(default):
            info, samples = read_input(default)
            if samples:
                jobs.append(default)
    if not jobs:
        print("Không có dữ liệu: hàng chờ HANG_CHO trống.")
        return 2
    all_created, rc = [], 0
    for path in jobs:
        try:
            info, samples = read_json(path) if path.lower().endswith(".json") else read_input(path)
            code, created = process(info, samples, os.path.basename(path), path)
        except Exception as ex:  # 1 phiếu lỗi không chặn các phiếu khác
            print(f"\n*** Lỗi khi xử lý {os.path.basename(path)}: {ex}")
            code, created = 2, []
        rc = max(rc, code)
        all_created += created
        blank = os.path.join(ROOT, "scripts", "_NHAP_MAU_TRONG.xlsx")
        if code == 0 and os.path.abspath(path) == os.path.join(ROOT, "NHAP_MAU.xlsx") and os.path.exists(blank):
            try:
                shutil.copy2(blank, path)
                print("   (NHAP_MAU.xlsx đã được làm trống cho lần sau — bản vừa nhập lưu ở DU_LIEU_NHAP)")
            except PermissionError:
                print("   ! NHAP_MAU.xlsx đang mở nên chưa làm trống — nhớ xóa dữ liệu cũ trước khi nhập phiếu mới")
        if code == 0 and os.path.dirname(os.path.abspath(path)) in queue_dirs():
            os.remove(path)  # đã lưu bản gốc ở DU_LIEU_NHAP — không giữ 2 nơi
    try:
        update_tool_zip()
    except Exception as ex:
        print("   ! Chưa cập nhật được PYC_TOOL.zip:", ex)
    with open(os.path.join(ROOT, "scripts", "_danh_sach_file_moi.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(all_created))
    return 0 if all_created else rc


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
